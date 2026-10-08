"""Virtual (online) public charter school enrollment in Oregon, 2019-20 on, from ODE's Fall Membership Reports
(October 1 counts by attending school). The 2025-26 report is the first to flag virtual schools; the virtual charter
schools are its rows with School Type "Charter School" and Virtual "Virtual", followed back by school ID through
earlier reports. ODE reports students by the school they attend, not the district they live in, so these are
statewide counts: how many Portland residents attend them is not published. Portland's figure is bounded by the
3% rule (ORS 338.125(2), OAR 581-020-0342): PPS says it is below 3% of its resident students.
Raw files (downloaded if missing): archive/ode-fall-membership/fallmembershipreport_<yyyyyyyy>.xlsx
Output: source/ode-virtual/virtual-charters.json
"""
import json, os, re, sys, urllib.request
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, 'archive', 'ode-fall-membership'); os.makedirs(RAW, exist_ok=True)
OUT = os.path.join(ROOT, 'source', 'ode-virtual'); os.makedirs(OUT, exist_ok=True)
URL = 'https://www.oregon.gov/ode/reports-and-data/students/Documents/fallmembershipreport_{}.xlsx'
YEARS = ['2019-20', '2020-21', '2021-22', '2022-23', '2023-24', '2024-25', '2025-26']
BANDS = {'K-5': ('Kindergarten', 'One', 'Two', 'Three', 'Four', 'Five'), '6-8': ('Six', 'Seven', 'Eight'),
         '9-12': ('Nine', 'Ten', 'Eleven', 'Twelve')}

def load(y):
    code = y[:4] + str(int(y[:4]) + 1)
    f = os.path.join(RAW, f'fallmembershipreport_{code}.xlsx')
    if not os.path.exists(f): urllib.request.urlretrieve(URL.format(code), f)
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    sheet = lambda p: next(w for w in wb.worksheets if w.title.strip().startswith(p))
    rows = list(sheet('School').iter_rows(values_only=True)); h = [str(c or '').strip() for c in rows[0]]
    col = lambda pat: next(i for i, c in enumerate(h) if re.search(pat, c, re.I))
    sid, sname = col(r'school.*(institution )?id'), col(r'^school( name)?$')
    tot = [i for i, c in enumerate(h) if 'Total Enrollment' in c][-1]
    pre = h[tot][:4]   # this year's columns start with its first year, e.g. "2025-26" or "2023-2024"
    grade = {g: next(i for i, c in enumerate(h) if c.startswith(pre) and re.search(rf'\b(grade )?{g}\b', c, re.I) and '%' not in c)
             for b in BANDS.values() for g in b}
    d = list(sheet('District').iter_rows(values_only=True)); dh = [str(c or '').strip() for c in d[0]]
    dn, dt = dh.index('District Name'), [i for i, c in enumerate(dh) if 'Total Enrollment' in c][-1]
    num = lambda v: v if isinstance(v, (int, float)) else 0
    state = sum(num(r[dt]) for r in d[1:] if r[dn])
    pps = next(num(r[dt]) for r in d[1:] if r[dn] == 'Portland SD 1J')
    return h, rows[1:], sid, sname, tot, grade, state, pps, num

h, rows, *_ = load('2025-26')
vi, ti = h.index('Virtual'), h.index('School Type')
VC = {r[1]: dict(name=r[2], sponsor=r[0]) for r in rows if r[vi] == 'Virtual' and r[ti] == 'Charter School'}
out = dict(schools={k: dict(v, enrollment={}) for k, v in VC.items()}, total={}, bands={}, state_total={}, pps_total={})
for y in YEARS:
    h, rows, sid, sname, tot, grade, state, pps, num = load(y)
    t = 0; bands = dict.fromkeys(BANDS, 0)
    for r in rows:
        if r[sid] not in VC or not isinstance(r[tot], (int, float)): continue
        out['schools'][r[sid]]['enrollment'][y] = r[tot]; t += r[tot]
        for b, gs in BANDS.items(): bands[b] += sum(num(r[grade[g]]) for g in gs)
    out['total'][y], out['bands'][y], out['state_total'][y], out['pps_total'][y] = t, bands, state, pps
    print(y, f'virtual charters {t:,} of {state:,} statewide ({100 * t / state:.1f}%), PPS {pps:,}', bands)
out['schools'] = sorted(out['schools'].values(), key=lambda s: -s['enrollment'].get('2025-26', 0))
J = dict(source="Oregon Department of Education, Fall Membership Reports 2019-20 to 2025-26 (school tab; virtual flag from 2025-26)",
         source_url='https://www.oregon.gov/ode/reports-and-data/students/Pages/Student-Enrollment-Reports.aspx',
         pps_cap=dict(quote='Portland Public Schools is currently below the 3% district cap set by OAR 581-020-0342.',
                      url='https://www.pps.net/Page/1029', retrieved='2026-10-08', note='undated on the PPS page'),
         years=YEARS, **out)
json.dump(J, open(os.path.join(OUT, 'virtual-charters.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('wrote', len(J['schools']), 'schools')
